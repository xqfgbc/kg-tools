"""
PPTX 结构化提取脚本 (python-pptx 通道)
=======================================
提取: 文本 + 坐标(x,y,w,h) + 形状类型 + GROUP递归 + 连接线 + 表格。

用法:
    python3 extract_pptx.py input.pptx > extraction_raw.txt 2>&1

关键: 内置 GROUP 递归，避免遗漏嵌套内容。
"""

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
import sys

pptx_path = sys.argv[1] if len(sys.argv) > 1 else "input.pptx"

prs = Presentation(pptx_path)

print(f"SLIDE_COUNT: {len(prs.slides)}")
print(f"SLIDE_WIDTH: {prs.slide_width} EMU = {prs.slide_width/914400:.1f} inches")
print(f"SLIDE_HEIGHT: {prs.slide_height} EMU = {prs.slide_height/914400:.1f} inches")
print("=" * 80)


def extract_text_recursive(shape, depth=0, max_depth=5):
    """
    递归提取形状中所有文本，含 GROUP 内部子形状。
    这是整个提取过程中最关键的函数——GROUP 内的文本不会被顶层遍历捕获。
    """
    results = []
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            text = para.text.strip()
            if text:
                results.append((depth, shape.name, text))
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP and depth < max_depth:
        try:
            for child in shape.shapes:
                results.extend(extract_text_recursive(child, depth + 1, max_depth))
        except:
            pass
    return results


for i, slide in enumerate(prs.slides):
    print(f"\n{'='*80}")
    print(f"SLIDE {i+1}")
    print(f"Layout: {slide.slide_layout.name}")
    print(f"Shapes count: {len(slide.shapes)}")
    print(f"{'='*80}")

    for j, shape in enumerate(slide.shapes):
        x_in = shape.left / 914400 if shape.left else 0
        y_in = shape.top / 914400 if shape.top else 0
        w_in = shape.width / 914400 if shape.width else 0
        h_in = shape.height / 914400 if shape.height else 0
        shape_type_str = str(shape.shape_type) if shape.shape_type else "UNKNOWN"

        # 跳过不可见的 think-cell 占位符
        if 'think-cell' in shape.name.lower():
            continue

        print(f"\n  Shape {j+1}: name='{shape.name}' type={shape_type_str}")
        print(f"    Position: left={x_in:.2f}in top={y_in:.2f}in width={w_in:.2f}in height={h_in:.2f}in")

        # 递归提取文本
        all_texts = extract_text_recursive(shape)
        if all_texts:
            for depth, sname, text in all_texts:
                indent = "    " + "  " * depth
                if depth == 0:
                    print(f"    TEXT: {text}")
                else:
                    print(f"    {indent}[{sname}] {text}")

        # 形状类型标记
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                print(f"    GROUP with {len(shape.shapes)} sub-shapes")
            except:
                pass
        elif shape.shape_type == MSO_SHAPE_TYPE.LINE:
            print(f"    LINE/CONNECTOR")
        elif shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            try:
                print(f"    TABLE: rows={len(shape.table.rows)}")
                for ri, row in enumerate(shape.table.rows):
                    row_texts = [cell.text.strip().replace('\n', ' | ') for cell in row.cells]
                    print(f"    ROW {ri}: {' || '.join(row_texts)}")
            except:
                pass

print("\n\nDONE")
